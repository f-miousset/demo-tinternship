"""Build the demo's seed data by running the fictional world through the real backend.

    make seed     # = cd app/backend && uv run python ../../scripts/build_seed.py

No model is called and nothing touches the network. What *is* used is the real
backend's own code — the base-document reader, the blank finder, the letter's
date and employer resolver, the preview renderer, the .docx writer, the job and
artifact serialisers and the rubrics — against an empty database in a temporary
directory. That is the point: the browser mock then serves bodies the backend
itself produced, instead of a hand-typed imitation that drifts from it on the
next resync (documentation/syncing-from-v2.md).

Outputs, all committed:

* `demo/static/seed/seed.json` — every table the mock needs, already serialised;
  `canned.json` beside it — every document a demo run can "write".
* `demo/static/artifacts/<id>.html` / `.docx` — the preview and the download for
  every résumé and letter a demo run can "write", since an iframe and a download
  link are not `fetch` calls the mock can answer.
* `demo/static/export-names.json` — the filename each download is served under.
* `selfhost/example-documents/` — Alex's four base documents, for trying the
  real app (selfhost/README.md) without writing a CV first.

Deterministic: the same world produces the same bytes, so a rebuild with no
change to `seed/world.py` or `app/` leaves `git status` clean.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import shutil
import sys
import tempfile
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "seed"))

DATA = Path(tempfile.mkdtemp(prefix="demo-seed-"))
os.environ.update(
    DATA_DIR=str(DATA),
    GOOGLE_API_KEY="demo-seed-never-used",
    BROWSER_VERIFY="false",
    FOLLOW_UP_AFTER_DAYS="14",
)

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

import world  # noqa: E402
from tinternship_backend.agents import rubrics  # noqa: E402
from tinternship_backend.agents.critic import weighted_overall  # noqa: E402
from tinternship_backend.agents.schemas import MasterProfile  # noqa: E402
from tinternship_backend.config import get_settings  # noqa: E402
from tinternship_backend.observability import pricing  # noqa: E402

get_settings.cache_clear()

from sqlmodel import select  # noqa: E402

from tinternship_backend.api.jobs import serialise_job  # noqa: E402
from tinternship_backend.db import engine as engine_module  # noqa: E402
from tinternship_backend.db.engine import session_scope  # noqa: E402
from tinternship_backend.db.models import (  # noqa: E402
    AgentRun,
    Application,
    ApplicationArtifact,
    ApplicationEvent,
    AuditRecord,
    JobPosting,
    PromptKind,
    TraceEvent,
)
from tinternship_backend.services import (  # noqa: E402
    base_documents,
    candidate_lists,
    docx_template,
    flows,
    letter_fields,
    profile_store,
    prompt_store,
    render,
)

engine_module._engine = None
engine_module.init_db()

EPOCH = datetime.fromisoformat(world.EPOCH)
TODAY = EPOCH.date()
MODELS = {"primary": "gemini-3.6-flash", "fast": "gemini-3.5-flash-lite", "grounded": "gemini-3.6-flash"}
OUT_STATIC = ROOT / "demo" / "static"
OUT_SEED = OUT_STATIC / "seed" / "seed.json"
OUT_EXAMPLES = ROOT / "selfhost" / "example-documents"
CANNED_BASE = 1000  # canned artifact ids start here; seeded ones are small integers


def normalise_docx(data: bytes) -> bytes:
    """The same document as the same bytes: python-docx stamps the save time
    into `docProps/core.xml` and the zip entries, which would make every rebuild
    a diff of 114 binary files."""
    import re
    import zipfile

    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            body = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)", rb"\g<1>2026-09-01T00:00:00Z\g<2>", body)
            entry = zipfile.ZipInfo(info.filename, date_time=(2026, 9, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, body)
    return buffer.getvalue()


def ago(days: float, hours: float = 0) -> datetime:
    return EPOCH - timedelta(days=days, hours=hours)


def stable(*parts: object) -> float:
    """A deterministic number in [0, 1) — the seed's only source of variation."""
    digest = hashlib.sha256("|".join(map(str, parts)).encode()).hexdigest()
    return int(digest[:8], 16) / 0x100000000


# ---------------------------------------------------------------------------
# Base documents — Alex's own CV and letter, with the blanks a run fills
# ---------------------------------------------------------------------------

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def _para(document, text: str = "", *, bold: bool = False, size: float = 10, center: bool = False,
          color: RGBColor | None = None, space_after: float = 2):
    paragraph = document.add_paragraph()
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(space_after)
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return paragraph


def _entry(document, left: str, right: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    head = paragraph.add_run(left)
    head.bold = True
    head.font.size = Pt(10)
    tail = paragraph.add_run(f"\t{right}")
    tail.font.size = Pt(10)


def _save(document) -> bytes:
    for section in document.sections:
        section.top_margin = section.bottom_margin = Pt(42)
        section.left_margin = section.right_margin = Pt(50)
    buffer = io.BytesIO()
    document.core_properties.author = "Alex Martin"
    document.core_properties.created = datetime(2026, 9, 1, tzinfo=UTC)
    document.core_properties.modified = datetime(2026, 9, 1, tzinfo=UTC)
    document.save(buffer)
    return buffer.getvalue()


RESUME_TEXT = {
    "en": {
        "contact": "alex.martin@example.org · +33 6 00 00 00 00 · Paris, France · EU citizen",
        "target": "[target_role] — six-month internship from February 2027",
        "profile": "[profile_summary]",
        "education": "EDUCATION",
        "msc": ("École Supérieure des Données (ESD) — MSc Data Science", "Paris · 2025–2027"),
        "msc_line": "Machine-learning track. Thesis: uncertainty in short-term load forecasting.",
        "courses": "Relevant coursework: [list_of_relevant_courses]",
        "bsc": ("Université Rive-Gauche — BSc Applied Mathematics, with honours", "Lyon · 2022–2025"),
        "experience": "EXPERIENCE",
        "job1": ("Brightwater Utilities — Data Analyst Intern", "Lyon · Jun–Aug 2025"),
        "job1_lines": [
            "Built a gradient-boosted daily demand forecast for 40 pumping stations (MAPE 6.1% → 4.3%).",
            "Automated the weekly leak-suspicion report in Python and SQL, saving about four hours a week.",
        ],
        "job2": ("ESD — Teaching Assistant, Statistics", "Paris · since Sep 2025"),
        "job2_lines": ["Weekly R and Python labs for 30 first-year students."],
        "projects": "PROJECTS",
        "proj1": "Clinical note de-identification — French token classifier, F1 0.94 (PyTorch, Transformers).",
        "proj2": "Bike-share rebalancing — MILP routing for 120 stations, −31% empty-station hours (OR-Tools).",
        "skills_h": "SKILLS",
        "skills": "Skills: [list_of_relevant_skills]",
        "languages": "Languages: French (native), English (C1), Spanish (B1)",
    },
    "fr": {
        "contact": "alex.martin@example.org · +33 6 00 00 00 00 · Paris · Nationalité française",
        "target": "[poste_vise] — stage de fin d'études de six mois à partir de février 2027",
        "profile": "[resume_du_profil]",
        "education": "FORMATION",
        "msc": ("École Supérieure des Données (ESD) — Master Data Science", "Paris · 2025–2027"),
        "msc_line": "Parcours machine learning. Mémoire : l'incertitude des prévisions de charge à court terme.",
        "courses": "Cours suivis : [liste_des_cours_pertinents]",
        "bsc": ("Université Rive-Gauche — Licence de mathématiques appliquées, mention bien", "Lyon · 2022–2025"),
        "experience": "EXPÉRIENCE",
        "job1": ("Brightwater Utilities — Stage Data Analyst", "Lyon · juin–août 2025"),
        "job1_lines": [
            "Prévision journalière de la demande de 40 stations de pompage par gradient boosting (MAPE 6,1 % → 4,3 %).",
            "Automatisation du rapport hebdomadaire de suspicion de fuites en Python et SQL (environ 4 h gagnées par semaine).",
        ],
        "job2": ("ESD — Chargé de TD, statistique", "Paris · depuis sept. 2025"),
        "job2_lines": ["TD hebdomadaires en R et Python pour 30 étudiants de première année."],
        "projects": "PROJETS",
        "proj1": "Pseudonymisation de notes cliniques — classifieur de tokens en français, F1 0,94 (PyTorch, Transformers).",
        "proj2": "Rééquilibrage de vélos partagés — tournées par MILP pour 120 stations, −31 % d'heures de stations vides (OR-Tools).",
        "skills_h": "COMPÉTENCES",
        "skills": "Compétences : [liste_des_competences]",
        "languages": "Langues : français (langue maternelle), anglais (C1), espagnol (B1)",
    },
}


def build_resume(language: str) -> bytes:
    t = RESUME_TEXT[language]
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)
    _para(document, "Alex Martin", bold=True, size=18, center=True, color=ACCENT, space_after=0)
    _para(document, t["contact"], size=9, center=True, space_after=6)
    _para(document, t["target"], bold=True, size=11, center=True, space_after=4)
    _para(document, t["profile"], space_after=8)
    for heading, body in (
        (t["education"], [("entry", t["msc"]), ("line", t["msc_line"]), ("line", t["courses"]), ("entry", t["bsc"])]),
        (t["experience"], [("entry", t["job1"]), *[("line", x) for x in t["job1_lines"]],
                           ("entry", t["job2"]), *[("line", x) for x in t["job2_lines"]]]),
        (t["projects"], [("line", t["proj1"]), ("line", t["proj2"])]),
        (t["skills_h"], [("line", t["skills"]), ("line", t["languages"])]),
    ):
        _para(document, heading, bold=True, size=11, color=ACCENT, space_after=2)
        for kind, value in body:
            if kind == "entry":
                _entry(document, *value)
            else:
                _para(document, value)
        _para(document, "", space_after=2)
    return _save(document)


LETTER_TEXT = {
    "en": [
        "Alex Martin", "alex.martin@example.org · +33 6 00 00 00 00 · Paris", "",
        "Paris, [current_month_letters] [current_date], [current_year]", "",
        "[team_name]", "[company_name]", "[address_company]", "[city_company]", "",
        "Re: application for [job_posting_name]", "",
        "Dear Hiring Team,", "",
        "[cover_letter_opening_paragraph]", "",
        "[cover_letter_middle_paragraph]", "",
        "[cover_letter_closing_paragraph]", "",
        "I am available for six months from February 2027 under an end-of-studies internship agreement, and would be glad to talk.", "",
        "Kind regards,",
    ],
    "fr": [
        "Alex Martin", "alex.martin@example.org · +33 6 00 00 00 00 · Paris", "",
        "Paris, le [current_date] [current_month_letters] [current_year]", "",
        "[team_name]", "[company_name]", "[address_company]", "[city_company]", "",
        "Objet : candidature au poste de [job_posting_name]", "",
        "Madame, Monsieur,", "",
        "[cover_letter_opening_paragraph]", "",
        "[cover_letter_middle_paragraph]", "",
        "[cover_letter_closing_paragraph]", "",
        "Je suis disponible six mois à partir de février 2027, dans le cadre d'une convention de stage de fin d'études, et serais heureux d'échanger avec vous.", "",
        "Je vous prie d'agréer, Madame, Monsieur, l'expression de mes salutations distinguées.",
    ],
}


def build_letter(language: str) -> bytes:
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10.5)
    for index, line in enumerate(LETTER_TEXT[language]):
        _para(document, line, bold=index == 0, size=13 if index == 0 else 10.5,
              color=ACCENT if index == 0 else None, space_after=0 if line else 4)
    return _save(document)


BASE_FILES = {
    ("resume", "en"): "CV-Alex-Martin-EN.docx",
    ("resume", "fr"): "CV-Alex-Martin-FR.docx",
    ("cover_letter", "en"): "Cover-letter-Alex-Martin-EN.docx",
    ("cover_letter", "fr"): "Lettre-Alex-Martin-FR.docx",
}


def install_base_documents() -> dict[tuple[str, str], int]:
    ids: dict[tuple[str, str], int] = {}
    for (kind, language), name in BASE_FILES.items():
        data = build_resume(language) if kind == "resume" else build_letter(language)
        base_documents.validate(kind, data)
        path = base_documents.stored_path(kind, language, name)
        path.write_bytes(data)
        ids[(kind, language)] = base_documents.replace(kind, language, label=name, file_path=str(path)).id
        OUT_EXAMPLES.mkdir(parents=True, exist_ok=True)
        (OUT_EXAMPLES / name).write_bytes(normalise_docx(data))
    # Rows default to the wall clock; pin them so a rebuild is byte-identical.
    from tinternship_backend.db.models import ProfileSource

    with session_scope() as session:
        for source in session.exec(select(ProfileSource)).all():
            source.created_at = ago(27)
            session.add(source)
    return ids


# ---------------------------------------------------------------------------
# The fills a run would write — composed from each posting's own sentences
# ---------------------------------------------------------------------------


def picks(table, tags: set[str], language: str, limit: int) -> str:
    chosen = [en if language == "en" else fr for en, fr, item_tags in table if item_tags & (tags | {"all"})]
    return ", ".join(chosen[:limit])


def resume_fills(job: dict, language: str) -> dict[str, str]:
    title = job["title"] if language == "en" else job["title_fr"]
    lead = {
        "forecast": ("forecasting", "la prévision"), "energy": ("energy data", "les données énergétiques"),
        "nlp": ("NLP", "le TAL"), "or": ("optimisation", "l'optimisation"), "analytics": ("analytics", "l'analyse de données"),
        "data_eng": ("data pipelines", "les pipelines de données"), "vision": ("computer vision", "la vision par ordinateur"),
        "finance": ("quantitative modelling", "la modélisation quantitative"), "geo": ("spatial data", "les données spatiales"),
        "logistics": ("logistics", "la logistique"), "ml": ("machine learning", "le machine learning"),
    }
    focus = [lead[tag][0 if language == "en" else 1] for tag in sorted(job["tags"]) if tag in lead][:2]
    if language == "en":
        summary = (
            f"MSc Data Science student focused on {' and '.join(focus)}, applying for {job['company']}'s "
            f"{title.split(' — ')[0]} role. Shipped a demand forecast that cut error by a third and explained it to the people who used it."
        )
    else:
        summary = (
            f"Étudiant en Master Data Science, orienté vers {' et '.join(focus)}, candidat au poste de "
            f"{title.split(' — ')[0]} chez {job['company']}. A livré une prévision de demande qui a réduit l'erreur d'un tiers, et l'a expliquée à ses utilisateurs."
        )
    return {
        "target_role" if language == "en" else "poste_vise": title,
        "profile_summary" if language == "en" else "resume_du_profil": summary,
        "list_of_relevant_courses" if language == "en" else "liste_des_cours_pertinents": picks(world.COURSES, job["tags"], language, 4),
        "list_of_relevant_skills" if language == "en" else "liste_des_competences": picks(world.SKILLS, job["tags"], language, 7),
    }


def letter_fills(job: dict, language: str) -> dict[str, str]:
    title = job["title"] if language == "en" else job["title_fr"]
    why, proof = job["why"][0 if language == "en" else 1], job["proof"][0 if language == "en" else 1]
    city = job["location"].split(",")[0]
    if language == "en":
        return {
            "team_name": "Data Science Team",
            "address_company": "",
            "city_company": city if city != "Remote" else "",
            "job_posting_name": title,
            "cover_letter_opening_paragraph": f"I am applying for the {title} position at {job['company']} because {why}.",
            "cover_letter_middle_paragraph": (
                f"{proof}. That is the part of the work I want to do more of: building the model, then measuring "
                f"honestly where it fails. {job['keywords'][0][0].upper() + job['keywords'][0][1:]} and "
                f"{job['keywords'][1]} are both on the posting, and both are in what I have done rather than only in what I have studied."
            ),
            "cover_letter_closing_paragraph": (
                f"I would bring that habit of checking a model against the people who rely on it to {job['company']}, "
                "and I would like to learn how your team decides when a model is good enough to ship."
            ),
        }
    return {
        "team_name": "Équipe Data Science",
        "address_company": "",
        "city_company": city if city != "Remote" else "",
        "job_posting_name": title,
        "cover_letter_opening_paragraph": f"Je vous adresse ma candidature au poste de {title} chez {job['company']}, car {why}.",
        "cover_letter_middle_paragraph": (
            f"{proof}. C'est la partie du travail que je souhaite approfondir : construire le modèle, puis mesurer "
            "honnêtement là où il échoue. Mon Master et mon stage couvrent directement les compétences demandées."
        ),
        "cover_letter_closing_paragraph": (
            f"J'apporterais chez {job['company']} cette habitude de confronter un modèle à ceux qui s'en servent, "
            "et j'aimerais apprendre comment votre équipe décide qu'un modèle est prêt à être déployé."
        ),
    }


def document_content(kind: str, job: dict, language: str, job_row_id: int | None) -> dict[str, Any]:
    """An artifact's `content`, assembled exactly as `flows.applying_stream` saves one."""
    known = None
    if kind == "cover_letter":
        facts = letter_fields.LetterFacts(language=language, company=job["company"], today=TODAY)
        known = letter_fields.resolver(facts)
    plan = base_documents.plan(kind, language, known=known)
    wanted = resume_fills(job, language) if kind == "resume" else letter_fills(job, language)
    fills = []
    for slot in plan.slots:
        name = letter_fields.field_name(slot.token)
        if slot.key in plan.auto or name not in wanted:
            continue
        fills.append({"slot": slot.key, "text": wanted[name], "reason": "Written for this posting." if kind == "resume" else "Drafted from the posting and the profile."})
    noun = "CV" if language == "fr" and kind == "resume" else ("Lettre" if language == "fr" else ("CV" if kind == "resume" else "Cover letter"))
    content: dict[str, Any] = {
        "document_title": f"{noun} — Alex Martin — {job['title' if language == 'en' else 'title_fr']}, {job['company']}",
        "fills": fills,
        "changes_made": [
            "Named the role exactly as the posting does." if language == "en" else "Repris l'intitulé exact de l'offre.",
            "Chose the coursework and skills closest to the posting's requirements." if kind == "resume" and language == "en"
            else ("Choisi les cours et compétences les plus proches des exigences." if kind == "resume" else
                  ("Opened on something specific to this employer." if language == "en" else "Ouvert sur un élément propre à cet employeur.")),
        ],
        "language": language,
        "blocks": docx_template.blocks_payload(plan.blocks),
        "slots": docx_template.slots_payload(plan.slots),
        "base_document_id": None,
    }
    if kind == "resume":
        content["keywords_covered"] = [k for k in job["keywords"] if k.lower() in json.dumps(fills, ensure_ascii=False).lower()]
        content["keywords_missing"] = [k for k in job["keywords"] if k not in content["keywords_covered"]][:2]
        content["gap_analysis"] = job["risks"]
    else:
        content["facts_used"] = [job["proof"][0]]
        content["address_source"] = ""
    flows._record_known_fills(content, plan)
    return content


# ---------------------------------------------------------------------------
# Audits — the Critic's verdicts, on the real rubrics
# ---------------------------------------------------------------------------

JUSTIFY = {
    "truthfulness": "Every claim traces to the profile, the lists or the posting.",
    "selection": "The posting's own title, and the coursework closest to its requirements.",
    "keyword_coverage": "The posting's main terms appear where the profile supports them.",
    "fit": "Each fill reads as part of the candidate's own line.",
    "page_use": "Uses the room on each line without spilling the page.",
    "language": "Written throughout in the chosen language.",
    "specificity": "Names something this employer actually does.",
    "evidence": "One concrete, measured result rather than adjectives.",
    "tone_and_language": "Plain register, no clichés.",
    "structure": "Each paragraph does one job.",
}


def verdict(kind: str, key: str, *, low: float = 7.6, high: float = 9.3, revision: int = 0) -> dict[str, Any]:
    rubric = rubrics.RUBRICS[kind]
    scores = [
        {
            "criterion": criterion.name,
            "score": round(low + (high - low) * stable(kind, key, criterion.name), 1),
            "justification": JUSTIFY.get(criterion.name, criterion.description.split(".")[0] + "."),
        }
        for criterion in rubric.criteria
    ]
    overall = round(weighted_overall({"scores": scores}, rubric), 2)
    return {
        "rubric": rubric.key,
        "scores": {"criteria": scores, "reasoning": "Clears the bar on every criterion; nothing invented."},
        "overall": overall,
        "threshold": 7.5,
        "verdict": "pass" if overall >= 7.5 else "revise",
        "blocking": rubric.blocking,
        "fixes": [],
        "revision": revision,
    }


# ---------------------------------------------------------------------------
# Runs and traces — what the Traces page shows
# ---------------------------------------------------------------------------

AGENTS = {
    "investigator": [("query_planner", "primary", []), ("scout", "grounded", ["google_search", "search_job_board"]),
                     ("normaliser", "fast", []), ("matcher", "primary", [])],
    "link_intake": [("reader", "fast", ["fetch_page"]), ("matcher", "primary", [])],
    "text_intake": [("reader", "fast", []), ("matcher", "primary", [])],
    "applying": [("resume_agent", "primary", []), ("address_scout", "grounded", ["google_search"]),
                 ("cover_letter_agent", "primary", []), ("cover_letter_humaniser", "fast", [])],
    "interview_prep": [("interview_researcher", "grounded", ["google_search"]), ("interview_coach", "primary", [])],
    "contacts": [("contact_scout", "grounded", ["google_search"]), ("contact_strategist", "primary", [])],
    "follow_up": [("follow_up_agent", "primary", [])],
    "interrogator": [("interrogator", "primary", [])],
    "hr_expert": [("hr_researcher", "grounded", ["google_search"]), ("playbook_writer", "primary", [])],
}


def add_run(kind: str, label: str, when: datetime, *, inp: dict, out: dict) -> int:
    """One run with a believable trace: agent start, model calls, tools, critic, end."""
    events: list[dict[str, Any]] = []
    t = when
    prompt_total = output_total = llm = tools = 0
    cost = 0.0

    def emit(**fields):
        nonlocal t
        t = t + timedelta(milliseconds=fields.get("latency_ms", 0) + 40)
        events.append({**fields, "created_at": t})

    emit(phase="run_start", agent_name=kind, summary=f"Run started: {label}", input_preview=json.dumps(inp, ensure_ascii=False)[:400])
    for name, tier, tool_names in AGENTS[kind]:
        model = MODELS[tier]
        emit(phase="agent_start", agent_name=name, model=model, summary=f"{name} started")
        for call in range(1 + (1 if tool_names else 0)):
            p = int(2400 + 6000 * stable(label, name, call))
            o = int(300 + 1500 * stable(label, name, call, "o"))
            latency = int(1800 + 5200 * stable(label, name, call, "l"))
            emit(phase="model_request", agent_name=name, model=model, summary=f"{name} → {model}",
                 input_preview="(instruction, context blocks, conversation so far)")
            emit(phase="model_response", agent_name=name, model=model, summary=f"{model}: {o} output tokens",
                 prompt_tokens=p, output_tokens=o, total_tokens=p + o, latency_ms=latency,
                 output_preview="(structured output, validated against the agent's schema)")
            prompt_total += p
            output_total += o
            llm += 1
            cost += pricing.estimate_cost_usd(model, p, o)
            if call == 0 and tool_names:
                for tool in tool_names:
                    emit(phase="tool_start", agent_name=name, tool_name=tool, summary=f"{tool} called",
                         input_preview=json.dumps({"query": label}, ensure_ascii=False))
                    emit(phase="tool_end", agent_name=name, tool_name=tool, summary=f"{tool} returned",
                         latency_ms=int(600 + 2400 * stable(label, tool)), output_preview="(results summarised for the next agent)")
                    tools += 1
        emit(phase="agent_end", agent_name=name, summary=f"{name} finished")
    emit(phase="run_end", agent_name=kind, summary="Run finished", output_preview=json.dumps(out, ensure_ascii=False)[:400])

    with session_scope() as session:
        run = AgentRun(
            kind=kind, label=label, status="succeeded", session_id=f"demo-{kind}", invocation_id=f"e-{hashlib.sha1(label.encode()).hexdigest()[:12]}",
            input=inp, output=out, prompt_tokens=prompt_total, output_tokens=output_total,
            total_tokens=prompt_total + output_total, llm_calls=llm, tool_calls=tools, cost_usd=round(cost, 5),
            started_at=when, finished_at=t,
        )
        session.add(run)
        session.flush()
        for seq, event in enumerate(events):
            session.add(TraceEvent(run_id=run.id, invocation_id=run.invocation_id, session_id=run.session_id, seq=seq,
                                   payload={}, **event))
        return int(run.id)


def add_audit(kind: str, subject_id: str, run_id: int | None, when: datetime, key: str, **kw) -> None:
    v = verdict(kind, key, **kw)
    with session_scope() as session:
        session.add(AuditRecord(subject_kind=rubrics.RUBRICS[kind].subject_kind, subject_id=subject_id, run_id=run_id,
                                created_at=when, **v))


# ---------------------------------------------------------------------------
# Interview brief, contacts, follow-up — read in the app, no file
# ---------------------------------------------------------------------------


def interview_prep(job: dict) -> dict[str, Any]:
    company = job["company"]
    return {
        "company_brief": f"{company} — {job['summary']} (Illustrative research: in the real app every line here is backed by a source the researcher opened.)",
        "recent_developments": [f"{company} published a hiring page for its data team this quarter.", "Expanded the team behind this role from three to four people."],
        "role_brief": job["description"].split("\n\n")[-1],
        "interview_process": ["Recruiter call (20 min)", "Team interview with a project walk-through (45 min)", "Take-home or live exercise", "Meeting with the manager"],
        "pitch": f"I'm finishing an MSc in Data Science at ESD. Last summer I built a demand forecast for a water utility that cut error by a third — and the part I enjoyed most was explaining its failure days to the operations team. {company}'s work on {job['keywords'][0]} is the same problem with higher stakes, which is why I applied.",
        "strengths_to_lead_with": job["strengths"],
        "gaps": [{"gap": risk, "severity": "medium", "honest_answer": "Say so plainly, then bridge to the closest thing you have done.", "preparation": "Read one recent public write-up on the topic and have a question ready about it."} for risk in job["risks"][:2]],
        "likely_questions": [
            {"question": "Walk us through your forecasting project — what would you do differently?", "asked_by": "Team lead", "why_asked": "Checks depth and self-criticism.",
             "answer_outline": ["Problem and baseline", "What moved the error", "The failure days and why", "What I'd change: quantile forecasts"], "evidence": ["Brightwater: MAPE 6.1% → 4.3%"]},
            {"question": f"Why {company} rather than a larger company?", "asked_by": "Manager", "why_asked": "Motivation.",
             "answer_outline": ["The specific problem", "Size of the team and ownership", "What I want to learn"], "evidence": [job["why"][0]]},
            {"question": "How would you explain a model's uncertainty to a non-technical colleague?", "asked_by": "Manager", "why_asked": "Communication.",
             "answer_outline": ["Use ranges, not a number", "Show a day it was wrong", "Say what decision changes"], "evidence": ["Weekly review with the operations team"]},
        ],
        "questions_to_ask": ["How do you decide a model is good enough to replace the current one?", "What happened to the last intern's project?", "Who uses this model's output day to day?"],
        "technical_topics": job["keywords"][:5],
        "logistics": ["Bring the convention de stage dates", "Have the Brightwater forecast plot ready to share"],
        "red_flags_to_probe": ["Whether the role is really modelling or mostly reporting"],
        "sources": [{"title": f"{company} careers page (illustrative)", "url": job["url"] or "https://careers.example.com/", "takeaway": "Team size and the role's missions."}],
        "language": "en",
    }


def contacts(job: dict) -> dict[str, Any]:
    company = job["company"]
    slug = company.lower().replace(" ", "-").replace("é", "e")
    people = []
    for index, (name, role, category) in enumerate(world.PEOPLE):
        people.append({
            "name": name, "role": role, "category": category,
            "why": {
                "hiring_manager": "Leads the team this role sits in — the person who will read the take-home.",
                "team_member": "Wrote the team's public post on its forecasting stack.",
                "recruiter": "Runs the internship campaign named on the posting.",
                "alumni": "ESD graduate — the shared school is the easiest opening.",
            }[category],
            "opening_line": "", "confidence": "medium" if index else "high",
            "linkedin_url": "", "evidence_url": f"https://careers.example.com/{slug}/team",
            "profile_status": "unproved", "evidence_status": "ok",
            "search_url": f"https://www.example.com/search/people?keywords={name.replace(' ', '%20')}%20{company.replace(' ', '%20')}",
            "web_url": f"https://www.example.com/search?q={name.replace(' ', '+')}+{company.replace(' ', '+')}",
        })
    return {
        "people": people,
        "dropped": [{"name": "J. Dupuis", "role": "Former intern", "reason": "Named on a 2023 page; no evidence they are still there."}],
        "angles": [
            {"key": "team", "label": "The data team", "why": "Who the role reports to.", "keywords": f'"{company}" "data scientist"',
             "url": f"https://www.example.com/search/people?keywords=%22{company.replace(' ', '%20')}%22%20data"},
            {"key": "alumni", "label": "ESD alumni there", "why": "A shared school gets answered.", "keywords": f'"{company}" "École Supérieure des Données"',
             "url": f"https://www.example.com/search/people?keywords=%22{company.replace(' ', '%20')}%22%20ESD"},
        ],
        "company": {"slug": slug, "url": f"https://www.example.com/company/{slug}"},
        "approach": ["Write to the team lead first, with one sentence on their published work.", "If no answer in a week, ask the alumna for a 15-minute call."],
        "notes": "Illustrative shortlist: every person here is fictional. In the real app each name was read on a page and each profile link was opened and proved.",
        "sources": [{"title": f"{company} team page (illustrative)", "url": f"https://careers.example.com/{slug}/team", "takeaway": "Names and roles of the data team."}],
        "checked": {"people_named": 5, "profiles_verified": 0, "dropped": 1, "searches": 2},
        "language": "en",
    }


def contact_message(job: dict, name: str) -> str:
    first = name.split()[0]
    return (f"Hello {first} — I'm an MSc Data Science student at ESD applying for {job['company']}'s {job['title'].split(' — ')[0]} internship. "
            f"Your team's work on {job['keywords'][0]} is exactly what I'd like to work on; could I ask you two questions about it?")


def follow_up(job: dict, applied_days: int) -> dict[str, Any]:
    paragraphs = [
        f"I applied for the {job['title']} internship on the careers page {applied_days} days ago and wanted to check where the process stands.",
        "Since applying I have finished the probabilistic-forecasting module of my MSc, which is directly relevant to the pricing simulations in the posting.",
        "Could you tell me whether the application is still under review, and roughly when I might hear back?",
    ]
    body = " ".join(paragraphs)
    return {
        "recipient": "the recruiting team", "to_hint": f"Reply to the confirmation e-mail from {job['company']}'s careers page, or write to the recruiter named on the posting.",
        "subject": f"Application for {job['title']} — Alex Martin", "greeting": "Hello,",
        "paragraphs": paragraphs, "sign_off": "Kind regards,", "signature": "Alex Martin · +33 6 00 00 00 00",
        "language": "en", "word_count": len(body.split()),
        "facts_used": [f"Applied {applied_days} days ago", "MSc module finished"],
        "send_notes": ["The posting names no recruiter — reply to the confirmation e-mail so it lands in the same thread."],
    }


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------


def main() -> None:
    if OUT_STATIC.exists():
        shutil.rmtree(OUT_STATIC)
    (OUT_STATIC / "artifacts").mkdir(parents=True)

    base_ids = install_base_documents()
    profile_store.save_profile(MasterProfile(**world.PROFILE))
    for kind, table in (("course", world.COURSES), ("skill", world.SKILLS)):
        for en, fr, _tags in table:
            candidate_lists.add(kind, en, fr)

    # Steps 2–4 of Account: the interview, the brief, the playbook.
    interview_run = add_run("interrogator", "Interview — internship preferences", ago(26), inp={"turns": len(world.INTERVIEW) // 2}, out={"brief": world.BRIEF["headline"]})
    brief = prompt_store.save_prompt(PromptKind.SEARCH_BRIEF, content=prompt_store.render_search_brief(world.BRIEF),
                                     structured=world.BRIEF, title=world.BRIEF["headline"], author="agent",
                                     note="Written by the brief writer from the interview.", invocation_id="demo-brief")
    add_audit("brief", str(brief.id), interview_run, ago(26), "brief")
    hr_run = add_run("hr_expert", "HR Expert — data science internships, France", ago(26, -1), inp={"brief": world.BRIEF["headline"]}, out={"sections": 12})
    playbook = prompt_store.save_prompt(PromptKind.PLAYBOOK, content=prompt_store.render_playbook(world.PLAYBOOK),
                                        structured=world.PLAYBOOK, title="Hiring playbook — data science internships, France",
                                        author="agent", note="Researched by the HR Expert.", invocation_id="demo-playbook")
    add_audit("playbook", str(playbook.id), hr_run, ago(26, -1), "playbook")
    from tinternship_backend.db.models import AgentPrompt

    with session_scope() as session:
        for prompt in session.exec(select(AgentPrompt)).all():
            prompt.created_at = ago(26) if prompt.kind == PromptKind.SEARCH_BRIEF else ago(26, -1)
            session.add(prompt)

    # Postings that already exist, each tied to the run that found it.
    rows: dict[str, int] = {}
    search_runs: dict[int, int] = {}
    for job in world.JOBS:
        if job["pool"] in ("link", "text") or job["pool"].startswith("search:"):
            continue
        day = job["discovered_days_ago"]
        if day not in search_runs:
            search_runs[day] = add_run("investigator", f"Search — {TODAY - timedelta(days=day):%d %b}", ago(day, -2),
                                       inp={"target_results": 15, "focus": ""}, out={"saved": 3})
            add_audit("job_ranking", "", search_runs[day], ago(day, -2), f"rank{day}", low=7.4, high=9.0)
        with session_scope() as session:
            row = JobPosting(
                dedupe_key=job["url"] or job["key"], **{k: job[k] for k in (
                    "title", "company", "location", "remote", "url", "source", "description", "summary", "requirements",
                    "nice_to_have", "contract_type", "start_date", "duration", "compensation", "language", "deadline",
                    "fit_score", "fit_rationale", "strengths", "risks", "keywords", "confidence", "url_status", "url_http_status")},
                apply_url=job["url"], posted_at=f"{job['posted_days_ago']} days ago", posted_on=(TODAY - timedelta(days=job["posted_days_ago"])).isoformat(),
                dismissed=job["pool"] == "discarded", run_id=search_runs[day],
                deferred_at=ago(job["deferred_days_ago"]) if job.get("deferred_days_ago") else None,
                discovered_at=ago(day, -2),
            )
            session.add(row)
            session.flush()
            rows[job["key"]] = int(row.id)

    # Applications, their timelines and the documents already written for them.
    jobs_by_key = {job["key"]: job for job in world.JOBS}
    app_ids: dict[str, int] = {}
    for key, spec in world.APPLICATIONS.items():
        job = jobs_by_key[key]
        first = spec["events"][0][0]
        with session_scope() as session:
            application = Application(
                job_posting_id=rows[key], status=spec["status"], pinned=spec["pinned"], language=spec["language"],
                notes=spec["notes"], next_action=spec["next_action"],
                next_action_date=(TODAY + timedelta(days=spec["next_action_days_ahead"])).isoformat() if spec.get("next_action_days_ahead") else "",
                applied_at=ago(spec["applied_days_ago"]) if spec.get("applied_days_ago") else None,
                trashed_at=ago(spec["trashed_days_ago"]) if spec.get("trashed_days_ago") else None,
                personalisation=world.PERSONALISATION["en"] if spec["documents"] else "",
                created_at=ago(first), updated_at=ago(spec["events"][-1][0]),
            )
            session.add(application)
            session.flush()
            app_ids[key] = int(application.id)
            for days, before, after, note in spec["events"]:
                session.add(ApplicationEvent(application_id=application.id, from_status=before, to_status=after, note=note, created_at=ago(days)))

        def save(kind: str, content: dict, run_id: int, when: datetime, revisions: int = 0, key: str = key) -> int:
            with session_scope() as session:
                artifact = ApplicationArtifact(application_id=app_ids[key], kind=kind, version=1, content=content,
                                               run_id=run_id, invocation_id=f"demo-{kind}-{key}", revisions=revisions, created_at=when)
                session.add(artifact)
                session.flush()
                return int(artifact.id)

        for language in spec["documents"]:
            when = ago(next(d for d, _b, after, _n in spec["events"] if after in ("preparing", "applied")) + 0.01)
            run = add_run("applying", f"Applying — {job['company']}", when, inp={"language": language, "resume": True, "cover_letter": True},
                          out={"artifacts": ["resume", "cover_letter"]})
            for kind in ("resume", "cover_letter"):
                revisions = 1 if stable(key, kind) > 0.6 else 0
                save(kind, document_content(kind, job, language, rows[key]), run, when, revisions)
                if revisions:
                    add_audit(kind, "", run, when, f"{key}{kind}r0", low=6.2, high=7.6, revision=0)
                add_audit(kind, "", run, when + timedelta(seconds=1), f"{key}{kind}", revision=revisions)
        if spec.get("interview_prep"):
            when = ago(7)
            run = add_run("interview_prep", f"Interview brief — {job['company']}", when, inp={"status": "hr_interview"}, out={"sources": 1})
            save("interview_prep", interview_prep(job), run, when)
            add_audit("interview_prep", "", run, when, f"{key}prep")
        if spec.get("contacts"):
            when = ago(6)
            run = add_run("contacts", f"Contacts — {job['company']}", when, inp={}, out={"people": 4})
            plan = contacts(job)
            plan["people"][0]["opening_line"] = contact_message(job, plan["people"][0]["name"])
            save("contacts", plan, run, when)
            add_audit("contacts", "", run, when, f"{key}contacts")
        if spec.get("follow_up"):
            when = ago(3)
            run = add_run("follow_up", f"Follow-up — {job['company']}", when, inp={"days_silent": spec["applied_days_ago"] - 3}, out={})
            save("follow_up", follow_up(job, spec["applied_days_ago"]), run, when)
            add_audit("follow_up", "", run, when, f"{key}fu")

    # Everything the mock can "write" later: résumé and letter for every posting
    # in both languages, interview briefs, shortlists and chase e-mails.
    canned: dict[str, Any] = {}
    names: dict[str, str] = {}
    next_id = CANNED_BASE
    for job in world.JOBS:
        entry: dict[str, Any] = {"documents": {}}
        for language in ("en", "fr"):
            for kind in ("resume", "cover_letter"):
                content = document_content(kind, job, language, rows.get(job["key"]))
                artifact_id = next_id
                next_id += 1
                (OUT_STATIC / "artifacts" / f"{artifact_id}.html").write_text(render.render_artifact(kind, content), encoding="utf-8")
                stem = render.document_stem(kind, candidate="Alex Martin", company=job["company"],
                                            title=job["title"] if language == "en" else job["title_fr"], language=language)
                path, _mime = flows._filled_docx(kind, content, stem=f"{stem}-{artifact_id}")
                (OUT_STATIC / "artifacts" / f"{artifact_id}.docx").write_bytes(normalise_docx(path.read_bytes()))
                names[str(artifact_id)] = f"{stem}.docx"
                entry["documents"][f"{kind}:{language}"] = {
                    "id": artifact_id, "content": content,
                    "text": render.artifact_text(kind, content, candidate_name="Alex Martin"),
                    "audit": {k: v for k, v in verdict(kind, f"canned{artifact_id}").items() if k not in ("blocking",)},
                }
        entry["interview_prep"] = interview_prep(job)
        entry["contacts"] = contacts(job)
        entry["contact_messages"] = {person[0]: contact_message(job, person[0]) for person in world.PEOPLE}
        entry["follow_up"] = follow_up(job, 15)
        canned[job["key"]] = entry

    # Serialise through the real API.
    from fastapi.testclient import TestClient

    from tinternship_backend.main import create_app

    client = TestClient(create_app())

    # The Settings page lists models by asking Google and the local Ollama.
    # Neither is asked here: the priced fallback list is what the page shows
    # when the listing fails, and it is the honest list for a demo with no key.
    from tinternship_backend.services import gemini_models, ollama_models

    gemini_models.available = lambda refresh=False: gemini_models._fallback()
    ollama_models.available = lambda *a, **k: []
    ollama_models.is_reachable = lambda *a, **k: False

    def get(path: str):
        response = client.get(path)
        response.raise_for_status()
        return response.json()

    applications = []
    for key, app_id in app_ids.items():
        detail = get(f"/api/applications/{app_id}")
        detail["artifacts_full"] = get(f"/api/applications/{app_id}/artifacts")["artifacts"]
        detail["job_key"] = key
        applications.append(detail)
    for application in applications:
        for artifact in application["artifacts_full"]:
            if artifact["kind"] in ("resume", "cover_letter"):
                text = client.get(f"/api/applications/artifacts/{artifact['id']}/text")
                artifact["text"] = text.text if text.status_code == 200 else ""
                (OUT_STATIC / "artifacts" / f"{artifact['id']}.html").write_text(
                    client.get(f"/api/applications/artifacts/{artifact['id']}/preview").text, encoding="utf-8")
                export = client.get(f"/api/applications/artifacts/{artifact['id']}/export")
                export.raise_for_status()
                (OUT_STATIC / "artifacts" / f"{artifact['id']}.docx").write_bytes(normalise_docx(export.content))
                disposition = export.headers.get("content-disposition", "")
                names[str(artifact["id"])] = disposition.split("filename=")[-1].strip('"') if "filename=" in disposition else f"document-{artifact['id']}.docx"

    all_jobs = {job["id"]: job for job in get("/api/jobs?view=all&limit=500")["jobs"]}
    key_of_row = {row_id: key for key, row_id in rows.items()}
    for row_id, job in all_jobs.items():
        job["key"] = key_of_row[row_id]

    # Postings revealed later, serialised by the same function as the rest so
    # the mock only has to give them an id when they appear.
    hidden = []
    for job in world.JOBS:
        if not (job["pool"] in ("link", "text") or job["pool"].startswith("search:")):
            continue
        row = JobPosting(
            id=0, dedupe_key=job["url"] or job["key"], **{k: job[k] for k in (
                "title", "company", "location", "remote", "url", "source", "description", "summary", "requirements",
                "nice_to_have", "contract_type", "start_date", "duration", "compensation", "language", "deadline",
                "fit_score", "fit_rationale", "strengths", "risks", "keywords", "confidence", "url_status", "url_http_status")},
            apply_url=job["url"], posted_at=f"{job['posted_days_ago']} days ago" if job["posted_days_ago"] else "today",
            posted_on=(TODAY - timedelta(days=job["posted_days_ago"])).isoformat(), discovered_at=EPOCH,
        )
        serialised = serialise_job(row, None)
        serialised.update(key=job["key"], pool=job["pool"])
        hidden.append(serialised)

    profile = get("/api/profile")
    seed = {
        "epoch": world.EPOCH,
        "version": "",
        "config": get("/api/config"),
        "account": get("/api/account"),
        "profile": profile,
        "lists": get("/api/profile/lists"),
        "prompts": get("/api/strategy/prompts")["prompts"],
        "settings_data": get("/api/settings/data"),
        "jobs": list(all_jobs.values()),
        "hidden_jobs": hidden,
        "applications": applications,
        "follow_ups": get("/api/applications/follow-ups"),
        "runs": [get(f"/api/traces/runs/{run['id']}") for run in get("/api/traces/runs?limit=500")["runs"]],
        "audits": get("/api/traces/audits?limit=500")["audits"],
        "settings_config": get("/api/settings/config"),
        "interview": [{"role": role, "author": "user" if role == "user" else "interrogator", "text": text} for role, text in world.INTERVIEW],
        "searches": world.SEARCHES,
        "link_import": world.LINK_IMPORT,
        "text_import": world.TEXT_IMPORT,
        "personalisation": world.PERSONALISATION,
        "base_files": {f"{kind}:{language}": name for (kind, language), name in BASE_FILES.items()},
        "base_ids": {f"{kind}:{language}": value for (kind, language), value in base_ids.items()},
        "messages": {
            # Ordered pairs, not objects: the file is written with sorted keys
            # for determinism, which would put the Matcher before the Planner.
            "investigator": list(flows.INVESTIGATOR_STEPS.items()), "link_intake": list(flows.LINK_INTAKE_STEPS.items()),
            "hr_expert": list(flows.HR_EXPERT_STEPS.items()),
            "stages": flows.STAGE_MESSAGES, "tracks": {str(k): v for k, v in flows.TRACK_LABELS.items()},
            "follow_up": flows.FOLLOW_UP_MESSAGES, "interview_prep": flows.INTERVIEW_PREP_MESSAGES, "contacts": flows.CONTACTS_MESSAGES,
        },
    }
    OUT_SEED.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(seed, ensure_ascii=False, indent=1, sort_keys=True, default=str)
    # The temporary data directory leaks into stored paths; nothing in the
    # browser needs them, and a path that changes per build defeats determinism.
    text = text.replace(str(DATA), "/data")
    # The version is the content's own hash: any rebuild that changes a byte
    # makes the browser start a returning visitor over (demo/src/mock/store.ts)
    # instead of mixing their saved state with a different world.
    text = text.replace('"version": ""', f'"version": "{hashlib.sha256(text.encode()).hexdigest()[:12]}"', 1)
    OUT_SEED.write_text(text + "\n", encoding="utf-8")
    # Every document a demo run can "write", in its own file: it is most of the
    # bytes, and the page loads it only when a run first needs it.
    canned_text = json.dumps(canned, ensure_ascii=False, sort_keys=True, default=sorted)
    (OUT_SEED.parent / "canned.json").write_text(canned_text.replace(str(DATA), "/data") + "\n", encoding="utf-8")
    (OUT_STATIC / "export-names.json").write_text(json.dumps(names, ensure_ascii=False, indent=1, sort_keys=True) + "\n", encoding="utf-8")
    shutil.rmtree(DATA, ignore_errors=True)
    print(f"seed: {len(all_jobs)} postings (+{len(hidden)} to reveal), {len(applications)} applications, "
          f"{len(seed['runs'])} runs, {len(names)} documents → {OUT_SEED.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
