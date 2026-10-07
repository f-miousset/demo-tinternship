"""Reading the documents a candidate can supply.

Three paths:

* **PDF** — handed to Gemini as an attachment rather than text-extracted. Designed
  and LinkedIn exports are graphic, multi-column layouts that `pypdf` turns to
  soup; the model reads the page as laid out. `pypdf` is still used for a cheap
  sanity check on page count and whether there is any text layer at all.
* **DOCX** — converted to text locally (python-docx), since the format is
  already structured and Gemini cannot ingest .docx directly.
* **LinkedIn data archive** — a ZIP of CSVs. Parsed deterministically; no model
  call, because the data is already structured and an LLM could only lose
  fidelity.
"""

from __future__ import annotations

import csv
import io
import logging
import zipfile
from dataclasses import dataclass, field
from typing import Any

from google.genai import types

from ..agents.schemas import (
    EducationEntry,
    ExperienceEntry,
    LanguageEntry,
    MasterProfile,
    ProjectEntry,
)

logger = logging.getLogger(__name__)

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass
class DocumentPayload:
    """What the extractor agent will actually receive."""

    text: str = ""
    parts: list[types.Part] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def pdf_payload(data: bytes, label: str) -> DocumentPayload:
    notes: list[str] = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = len(reader.pages)
        text_layer = "".join((page.extract_text() or "") for page in reader.pages[:3]).strip()
        notes.append(f"{pages} page(s)")
        if not text_layer:
            notes.append(
                "no extractable text layer — the document is image-based, so the model is "
                "reading it visually"
            )
    except Exception:
        logger.warning("Could not pre-inspect PDF %s", label, exc_info=True)

    return DocumentPayload(
        text=(
            f"The attached PDF is: {label}. Transcribe it into the structured profile. "
            f"Pre-check: {'; '.join(notes) if notes else 'n/a'}."
        ),
        parts=[types.Part.from_bytes(data=data, mime_type=PDF_MIME)],
        notes=notes,
    )


def docx_payload(data: bytes, label: str) -> DocumentPayload:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    body = "\n".join(lines)
    return DocumentPayload(
        text=(
            f"The following is the text of {label}, a Word document. Transcribe it into the "
            f"structured profile.\n\n---\n\n{body}"
        ),
        notes=[f"{len(lines)} text blocks"],
    )


def payload_for_upload(filename: str, data: bytes) -> DocumentPayload:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return pdf_payload(data, filename)
    if lower.endswith(".docx"):
        return docx_payload(data, filename)
    if lower.endswith((".txt", ".md")):
        return DocumentPayload(
            text=(
                f"The following is the content of {filename}. Transcribe it into the "
                f"structured profile.\n\n---\n\n{data.decode('utf-8', errors='replace')}"
            )
        )
    raise ValueError(
        f"Unsupported file type: {filename}. Upload a PDF, DOCX, TXT or a LinkedIn export ZIP."
    )


# ---------------------------------------------------------------------------
# LinkedIn data archive
# ---------------------------------------------------------------------------

_ARCHIVE_FILES = {
    "profile": ("Profile.csv",),
    "positions": ("Positions.csv",),
    "education": ("Education.csv",),
    "skills": ("Skills.csv",),
    "projects": ("Projects.csv",),
    "languages": ("Languages.csv",),
    "certifications": ("Certifications.csv",),
    "honors": ("Honors.csv", "Honors_and_Awards.csv"),
    "emails": ("Email Addresses.csv",),
}


def _read_csv(archive: zipfile.ZipFile, candidates: tuple[str, ...]) -> list[dict[str, str]]:
    names = {name.split("/")[-1]: name for name in archive.namelist()}
    for candidate in candidates:
        actual = names.get(candidate)
        if actual is None:
            continue
        try:
            with archive.open(actual) as handle:
                text = handle.read().decode("utf-8-sig", errors="replace")
            return list(csv.DictReader(io.StringIO(text)))
        except Exception:
            logger.warning("Failed to read %s from LinkedIn archive", candidate, exc_info=True)
    return []


def _get(row: dict[str, str], *keys: str) -> str:
    for key in keys:
        value = (row.get(key) or "").strip()
        if value:
            return value
    return ""


def parse_linkedin_archive(data: bytes) -> MasterProfile:
    """Deterministically build a profile from a LinkedIn data export ZIP."""
    profile = MasterProfile()
    notes: list[str] = []

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        rows = _read_csv(archive, _ARCHIVE_FILES["profile"])
        if rows:
            row = rows[0]
            first = _get(row, "First Name")
            last = _get(row, "Last Name")
            profile.full_name = " ".join(part for part in (first, last) if part)
            profile.headline = _get(row, "Headline")
            profile.summary = _get(row, "Summary")
            profile.location = _get(row, "Geo Location", "Location")
            website = _get(row, "Websites")
            if website:
                profile.portfolio_url = website
        else:
            notes.append("Profile.csv missing from the archive")

        emails = _read_csv(archive, _ARCHIVE_FILES["emails"])
        for row in emails:
            if _get(row, "Primary").lower() in {"yes", "true"} or not profile.email:
                profile.email = _get(row, "Email Address") or profile.email

        for row in _read_csv(archive, _ARCHIVE_FILES["positions"]):
            profile.experiences.append(
                ExperienceEntry(
                    title=_get(row, "Title"),
                    organisation=_get(row, "Company Name"),
                    location=_get(row, "Location"),
                    start_date=_get(row, "Started On"),
                    end_date=_get(row, "Finished On") or "Present",
                    description=_get(row, "Description"),
                )
            )

        for row in _read_csv(archive, _ARCHIVE_FILES["education"]):
            profile.education.append(
                EducationEntry(
                    degree=" ".join(
                        part
                        for part in (_get(row, "Degree Name"), _get(row, "Notes"))
                        if part
                    ).strip(),
                    institution=_get(row, "School Name"),
                    start_date=_get(row, "Start Date"),
                    end_date=_get(row, "End Date"),
                    details=_get(row, "Activities"),
                )
            )

        for row in _read_csv(archive, _ARCHIVE_FILES["projects"]):
            profile.projects.append(
                ProjectEntry(
                    name=_get(row, "Title"),
                    description=_get(row, "Description"),
                    url=_get(row, "Url"),
                )
            )

        profile.skills = [
            name for row in _read_csv(archive, _ARCHIVE_FILES["skills"]) if (name := _get(row, "Name"))
        ]
        profile.languages = [
            LanguageEntry(language=name, level=_get(row, "Proficiency"))
            for row in _read_csv(archive, _ARCHIVE_FILES["languages"])
            if (name := _get(row, "Name"))
        ]
        profile.certifications = [
            name
            for row in _read_csv(archive, _ARCHIVE_FILES["certifications"])
            if (name := _get(row, "Name"))
        ]
        profile.awards = [
            title for row in _read_csv(archive, _ARCHIVE_FILES["honors"]) if (title := _get(row, "Title"))
        ]

    if not profile.experiences and not profile.education:
        notes.append(
            "No positions or education found — check that this is the full LinkedIn data "
            "export and not the 'basic' one."
        )
    profile.extraction_notes = "; ".join(notes)
    return profile


def looks_like_linkedin_archive(filename: str, data: bytes) -> bool:
    if not filename.lower().endswith(".zip"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = {name.split("/")[-1] for name in archive.namelist()}
        return bool(names & {"Profile.csv", "Positions.csv", "Education.csv"})
    except Exception:
        return False


def summarise_payload(payload: DocumentPayload) -> dict[str, Any]:
    return {"notes": payload.notes, "has_attachment": bool(payload.parts)}
