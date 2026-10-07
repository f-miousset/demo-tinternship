"""One page means one page, and nothing here can lay out a Word document.

`services/page_fit.py` is the answer to the candidate's bluntest rule — *never
more than one page* — on a stack that has no renderer. It measures instead:
resolves each paragraph's size, indent and spacing, wraps its text against the
column with real Calibri advance widths, and adds it up.

That makes these tests two different things at once, and it is worth being
explicit about which is which:

* **Behavioural tests**, which are exact. Wider text wraps to more lines; a
  bigger font is taller; an indent narrows the column; a budget shrinks as a
  page fills. These hold whatever the width table says.
* **One calibration test**, which is not exact and says so. A document Word
  itself recorded as one page must not measure as two. That is the only ground
  truth available — `docProps/app.xml` carries `<Pages>`, written by Word at
  save time — and it is what keeps an estimator honest about the direction of
  its error.
"""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt

from tinternship_backend.services import docx_template as dt
from tinternship_backend.services import page_fit


def _docx(build) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _plain(lines: list[str], *, size: float = 11.0) -> bytes:
    def build(document):
        for line in lines:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.font.size = Pt(size)

    return _docx(build)


class TestWrapping:
    def test_a_short_line_is_one_line(self):
        assert page_fit.wrapped_lines([("Alex Martin", 11.0, False)], 400) == 1

    def test_a_long_line_wraps(self):
        text = "Ingénieur en apprentissage automatique " * 6

        assert page_fit.wrapped_lines([(text, 11.0, False)], 400) >= 3

    def test_the_same_text_wraps_more_in_a_narrower_column(self):
        text = "Conception d'un pipeline de données pour l'équipe analytique"
        wide = page_fit.wrapped_lines([(text, 11.0, False)], 400)
        narrow = page_fit.wrapped_lines([(text, 11.0, False)], 150)

        assert narrow > wide

    def test_the_same_text_wraps_more_at_a_bigger_size(self):
        text = "Conception d'un pipeline de données pour l'équipe analytique"

        assert page_fit.wrapped_lines([(text, 20.0, False)], 300) > page_fit.wrapped_lines(
            [(text, 9.0, False)], 300
        )

    def test_each_run_is_measured_at_its_own_size(self):
        """A heading line is regularly a big bold employer followed by a small
        location. Measuring the whole line at the larger of the two is how an
        estimate ends up wildly out."""
        mixed = [("Acme Corporation ", 24.0, True), ("Paris, France", 8.0, False)]
        uniform = [("Acme Corporation Paris, France", 24.0, True)]

        assert page_fit.wrapped_lines(mixed, 250) < page_fit.wrapped_lines(uniform, 250)

    def test_a_hard_break_starts_a_line(self):
        assert page_fit.wrapped_lines([("one\ntwo\nthree", 11.0, False)], 400) == 3

    def test_a_tabbed_header_fits_when_both_sides_fit(self):
        """`Employer<TAB>City` is one line exactly when the two pieces fit side
        by side, which is what a right tab stop does."""
        assert page_fit.wrapped_lines([("Acme\tParis, France", 11.0, False)], 400) == 1

    def test_a_word_too_long_for_the_column_is_broken_across_lines(self):
        """Word breaks it rather than letting it overhang. A wrapper that does
        not counts a 40,000-character run as one line — which is what a runaway
        fill looks like, and the exact case the page gate exists to refuse."""
        unbreakable = "x" * 4000

        assert page_fit.wrapped_lines([(unbreakable, 11.0, False)], 400) > 40

    def test_an_empty_line_still_takes_a_line(self):
        assert page_fit.wrapped_lines([("", 11.0, False)], 400) == 1


class TestMeasuring:
    def test_a_longer_document_is_taller(self):
        short = page_fit.measure(_plain(["one line"]))
        long = page_fit.measure(_plain([f"line {n}" for n in range(30)]))

        assert long.height > short.height

    def test_an_empty_document_fits(self):
        assert page_fit.measure(_plain([""])).fits

    def test_a_document_of_hundreds_of_lines_does_not(self):
        estimate = page_fit.measure(_plain([f"line {n}" for n in range(300)]))

        assert not estimate.fits
        assert estimate.fraction > 1

    def test_the_indent_of_a_bullet_narrows_its_column(self):
        """The indent lives in `numbering.xml`, not on the paragraph, so a
        reader that only looks at `w:ind` measures every bullet against a
        column that is 18pt too wide."""

        def build(document):
            document.add_paragraph("a" * 40, style="List Bullet")

        indented = page_fit.measure(_docx(build)).paragraphs[0]
        plain = page_fit.measure(_plain(["a" * 40])).paragraphs[0]

        assert indented.column < plain.column

    def test_space_after_counts_towards_the_height(self):
        def build(document):
            paragraph = document.add_paragraph("one line")
            paragraph.paragraph_format.space_after = Pt(40)

        assert page_fit.measure(_docx(build)).height > page_fit.measure(
            _plain(["one line"])
        ).height + 30

    def test_a_manual_page_break_is_more_than_one_page_however_it_is_filled(self):
        def build(document):
            document.add_paragraph("first page")
            document.add_page_break()
            document.add_paragraph("second page")

        estimate = page_fit.measure(_docx(build))

        assert estimate.page_break
        assert not estimate.fits


class TestSlackAndBudgets:
    def test_an_almost_empty_page_leaves_room(self):
        estimate = page_fit.measure(_plain(["one line"]))

        assert estimate.slack_lines > 40
        assert estimate.slack_chars > 1000

    def test_a_full_page_leaves_none(self):
        estimate = page_fit.measure(_plain([f"line {n}" for n in range(200)]))

        assert estimate.slack_chars == 0

    def test_a_budget_shrinks_as_the_page_fills(self):
        def document_with(filler: int) -> tuple:
            data = _plain(["Poste : [Poste visé]"] + [f"line {n}" for n in range(filler)])
            blocks = dt.read_blocks(data)
            slots = dt.find_slots(blocks)
            return page_fit.slot_budgets(page_fit.measure(data), slots)["s0"]

        assert document_with(5) > document_with(46) > document_with(47)

    def test_every_blank_can_always_hold_a_job_title(self):
        """Below the floor a blank cannot hold "Ingénieur Machine Learning",
        and a budget that small is not a budget, it is a refusal."""
        data = _plain(["Poste : [Poste visé]"] + [f"line {n}" for n in range(200)])
        blocks = dt.read_blocks(data)
        slots = dt.find_slots(blocks)

        budgets = page_fit.slot_budgets(page_fit.measure(data), slots)

        assert budgets["s0"] >= page_fit.SLOT_FLOOR

    def test_no_blank_may_become_a_paragraph(self):
        data = _plain(["Poste : [Poste visé]"])
        blocks = dt.read_blocks(data)
        slots = dt.find_slots(blocks)

        budgets = page_fit.slot_budgets(page_fit.measure(data), slots)

        assert budgets["s0"] <= page_fit.SLOT_CEILING + len("[Poste visé]")

    def test_required_room_is_what_the_blanks_cannot_do_without(self):
        blocks = dt.read_blocks(_plain(["Poste : [x]"]))

        assert page_fit.required_room(dt.find_slots(blocks)) == page_fit.SLOT_FLOOR - len("[x]")


class TestThePageGate:
    def _document(self, filler: int) -> tuple:
        data = _plain(["Poste : [Poste visé]"] + [f"line {n}" for n in range(filler)])
        blocks = dt.read_blocks(data)
        return data, blocks, dt.find_slots(blocks)

    def test_a_modest_fill_on_a_half_empty_page_passes(self):
        data, blocks, slots = self._document(5)

        assert page_fit.page_problems(data, blocks, slots, {"s0": "Ingénieur ML"}) == []

    def test_a_fill_that_spills_the_page_is_refused(self):
        data, blocks, slots = self._document(56)
        huge = "Ingénieur en apprentissage automatique et en optimisation " * 12

        problems = page_fit.page_problems(data, blocks, slots, {"s0": huge})

        assert any("second page" in problem for problem in problems)

    def test_one_enormous_unbreakable_fill_is_refused_too(self):
        data, blocks, slots = self._document(5)

        problems = page_fit.page_problems(data, blocks, slots, {"s0": "x" * 40000})

        assert any("second page" in problem for problem in problems)

    def test_the_refusal_names_the_longest_fill(self):
        """"Make it shorter" is not actionable; "s0 is 700 characters" is."""
        data, blocks, slots = self._document(56)
        huge = "mots " * 200

        problems = page_fit.page_problems(data, blocks, slots, {"s0": huge})

        assert any("s0 (" in problem for problem in problems)


class TestTheOtherHalfOfTheSameRule:
    """One page also means *a* page. Until 2026-09-10 only the ceiling was
    checked, and a model taught only where the cliff is camps well below it: the
    first cover letter filled this way used 1,131 of its 2,133 characters and
    stopped eight lines above the bottom of the page, and a résumé the same day
    listed six skills on a line with room for nine.

    The document these use is 96% full with two blanks and about 185 characters
    of room, which is the shape of the candidate's own CV (96%, four blanks, 268
    characters) rather than the near-empty page the other fixtures use. That is
    load-bearing: on a page with more room than its blanks could ever hold,
    this check correctly asks for nothing.
    """

    def _document(self, filler: int = 44, *, blanks: int = 2) -> tuple:
        lines = [f"Ligne {n} : [blanc {n}]" for n in range(blanks)]
        data = _plain(lines + [f"ligne numero {n}" for n in range(filler)])
        blocks = dt.read_blocks(data)
        slots = dt.find_slots(blocks)
        estimate = page_fit.measure(data)
        return estimate, dt.with_budgets(slots, page_fit.slot_budgets(estimate, slots))

    @staticmethod
    def _to_target(estimate, slots, alphabet: str = "x") -> dict[str, str]:
        """Fills that add exactly the target between them, split evenly."""
        share = -(-page_fit.total_target(estimate, slots) // len(slots))
        return {
            slot.key: (alphabet * 400)[: len(slot.token) + share] for slot in slots
        }

    def test_a_fill_that_uses_the_room_passes(self):
        estimate, slots = self._document()

        fills = self._to_target(estimate, slots)

        assert page_fit.short_problems(estimate, slots, fills) == []

    def test_a_fill_that_leaves_most_of_the_page_empty_is_refused(self):
        estimate, slots = self._document()

        problems = page_fit.short_problems(estimate, slots, {"s0": "ML", "s1": "Acme"})

        assert any("characters this page has room for" in problem for problem in problems)

    def test_the_refusal_names_the_longest_fill_and_its_budget(self):
        """The mirror of the page gate's message: which blank, and how much of
        its budget is still unspent."""
        estimate, slots = self._document()

        problems = page_fit.short_problems(estimate, slots, {"s0": "court", "s1": "un peu plus"})

        assert any("s1 (11 of " in problem for problem in problems)

    def test_a_page_with_no_room_left_asks_for_nothing(self):
        """The interlock that stops the two gates arguing. Below a line of free
        space "you left room unused" is noise, and asking for more of a page
        that has none is asking to be refused by `page_problems`."""
        estimate, slots = self._document(56)

        assert page_fit.short_problems(estimate, slots, {"s0": "ML", "s1": "Acme"}) == []

    def test_a_page_bigger_than_its_blanks_asks_for_nothing_either(self):
        """Two short blanks on a near-empty page. The room is real and the
        blanks cannot take it: `SLOT_CEILING` caps each of them at 400
        characters whatever the page says, so demanding they fill it is
        demanding a 400-character job title, in a loop with no winning move.
        Laying that page out is the candidate's job, not the model's."""
        estimate, slots = self._document(5)

        assert page_fit.total_target(estimate, slots) == 0
        assert page_fit.short_problems(estimate, slots, {"s0": "ML", "s1": "Acme"}) == []

    def test_there_is_a_band_between_the_two_gates(self):
        """`page_problems` refuses above the room, this refuses below 80% of it,
        so a fill written to the target passes both — which is the whole point
        of asking for 80% and not for all of it.

        With ordinary prose, which is what `slack_chars` is an average over. A
        fill made of nothing but wide characters spends more page per character
        than the document it was measured against and can still spill, which is
        why `hard_check` asks about the page first and only then about the
        target."""
        data = _plain(
            ["Ligne 0 : [blanc 0]", "Ligne 1 : [blanc 1]"]
            + [f"ligne numero {n}" for n in range(44)]
        )
        blocks = dt.read_blocks(data)
        slots = dt.find_slots(blocks)
        estimate = page_fit.measure(data)
        slots = dt.with_budgets(slots, page_fit.slot_budgets(estimate, slots))
        fills = self._to_target(
            estimate, slots, "gouvernance des donnees et fiabilisation des flux "
        )

        assert page_fit.short_problems(estimate, slots, fills) == []
        assert page_fit.page_problems(data, blocks, slots, fills) == []

    def test_an_address_left_empty_is_never_named_as_one_to_lengthen(self):
        """A street nobody published is correctly blank, and the one thing this
        must never do is ask for it to be filled in anyway."""
        estimate, slots = self._document()

        problems = page_fit.short_problems(
            estimate, slots, {"s0": "Pantin", "s1": ""}, optional={"s1"}
        )

        assert problems and "s1 (" not in problems[0]

    def test_what_python_already_answered_is_not_credited_to_the_model(self):
        """`added_chars` is given the open blanks only: a letter's date and
        employer are already in the document the room was measured on."""
        estimate, slots = self._document()

        both = page_fit.added_chars(slots, {"s0": "x" * 50, "s1": "x" * 50})
        one = page_fit.added_chars(slots[:1], {"s0": "x" * 50, "s1": "x" * 50})

        assert both > one


class TestWhatWordItselfRecorded:
    """The one place there is ground truth: Word writes `<Pages>` on save."""

    def test_the_page_count_is_read_when_word_wrote_one(self, tmp_path):
        import zipfile

        data = _plain(["one line"])
        patched = tmp_path / "patched.docx"
        with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(
            patched, "w"
        ) as target:
            for item in source.namelist():
                if item != "docProps/app.xml":
                    target.writestr(item, source.read(item))
            target.writestr(
                "docProps/app.xml",
                '<?xml version="1.0"?><Properties><Pages>3</Pages></Properties>',
            )

        assert page_fit.word_page_count(patched.read_bytes()) == 3

    def test_a_file_with_no_app_part_reports_nothing(self):
        import zipfile

        data = _plain(["one line"])
        stripped = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(
            stripped, "w"
        ) as target:
            for item in source.namelist():
                if item != "docProps/app.xml":
                    target.writestr(item, source.read(item))

        assert page_fit.word_page_count(stripped.getvalue()) == 0

    def test_python_docx_carries_its_templates_stale_number_forward(self):
        """Which is exactly why this is only ever read from the candidate's own
        upload. python-docx neither computes nor clears these statistics, so the
        default template's "one page" survives into a document of any length."""
        assert page_fit.word_page_count(_plain([f"line {n}" for n in range(400)])) == 1

    def test_something_that_is_not_a_docx_reports_nothing_rather_than_raising(self):
        assert page_fit.word_page_count(b"not a zip") == 0

    def test_word_saying_one_page_raises_the_ceiling_over_a_pessimistic_estimate(self):
        """The width table is a table, not a rasteriser. When Word recorded the
        candidate's own file as one page and this measured it at slightly more,
        Word is right — and holding the fills to a capacity their untouched
        document already "exceeds" would reject a no-op."""
        measured = page_fit.measure(_plain([f"line {n}" for n in range(120)]))
        fields = {
            "height": measured.height,
            "capacity": measured.capacity,
            "paragraphs": measured.paragraphs,
            "body_size": measured.body_size,
            "body_width": measured.body_width,
            "char_width": measured.char_width,
        }
        certified = page_fit.Estimate(**fields, word_pages=1)
        unknown = page_fit.Estimate(**fields, word_pages=0)

        assert not unknown.fits
        assert page_fit.allowance(certified) == certified.height
        assert page_fit.allowance(unknown) == unknown.capacity


class TestTheDirectionOfTheError:
    """The estimator is allowed to be a few percent out. It is not allowed to be
    out in the direction that refuses a document Word calls one page.

    That direction is the expensive one: an over-estimate blocks every French
    application the candidate ever makes, where an under-estimate costs one
    revision round in a gate that re-measures anyway. So the constants are set
    to leave a real one-page CV comfortably under the line, and the two things
    most likely to drift — the border model and the safety margin — are pinned
    here rather than left to be re-tuned by whoever next sees a number they
    dislike.
    """

    def test_a_page_of_ordinary_body_text_is_not_called_two(self):
        estimate = page_fit.measure(_plain([f"line {n}" for n in range(40)]))

        assert estimate.fits
        assert estimate.fraction < 1.0

    def test_a_rule_under_a_heading_costs_its_stroke_and_not_its_gap(self):
        """Counting `w:space` as well put the estimate a percent over a document
        Word itself records as one page. Five section rules is most of a line,
        so the difference decides pages."""

        def build(document):
            paragraph = document.add_paragraph("Education")
            borders = paragraph._p.get_or_add_pPr().makeelement(
                f"{page_fit.W}pBdr", {}
            )
            bottom = borders.makeelement(
                f"{page_fit.W}bottom",
                {f"{page_fit.W}sz": "6", f"{page_fit.W}space": "24"},
            )
            borders.append(bottom)
            paragraph._p.get_or_add_pPr().append(borders)
            return document

        ruled = page_fit.measure(_docx(build))
        plain = page_fit.measure(_plain(["Education"]))

        assert 0 < ruled.height - plain.height < 2


class TestDescribing:
    def test_it_quotes_a_percentage_the_candidate_can_disagree_with(self):
        said = page_fit.describe(page_fit.measure(_plain(["one line"])))

        assert "%" in said and "characters" in said

    def test_a_document_with_no_room_says_so_and_says_what_to_do(self):
        said = page_fit.describe(page_fit.measure(_plain([f"line {n}" for n in range(200)])))

        assert "Free up" in said

    def test_a_manual_page_break_is_reported_as_the_file_problem_it_is(self):
        def build(document):
            document.add_paragraph("first")
            document.add_page_break()
            document.add_paragraph("second")

        assert "page break" in page_fit.describe(page_fit.measure(_docx(build)))
